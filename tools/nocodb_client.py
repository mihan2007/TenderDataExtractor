#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
nocodb_client.py - syncs one extracted tender into NocoDB (step 7.1 in CLAUDE.md).

1. Upserts the tender into the registry table (matched by tender number).
2. With --xlsx, syncs the "Items & prices" sheet into the tender's OWN table
   (one table per tender, named by its internal number), creates the
   "Agent prices", "Procurement" and "Price comparison" views, and links the
   table from the registry record ("Specification link").
3. Re-reads everything it wrote and reports any field that did not land as
   intended ("verification" / "spec_verification" in the output).

Configuration (NocoDB URL, workspace/base/registry table IDs, base currency,
number of procurement slots) comes from config.json - see tools/config.py.

Usage:
    # first run on a new NocoDB base: create / complete the registry table
    python nocodb_client.py --init-registry --token-file secrets/API_key_NocoDB.txt

    # normal sync after save_tender.py
    python nocodb_client.py --json tender_data.json --token-file secrets/API_key_NocoDB.txt \
        --xlsx "<tender>/AIExtract/Items_and_prices.xlsx" [--rate USD=0.92 --rate CNY=0.13]

    # maintenance on an already saved tender (rows are not touched unless noted)
    python nocodb_client.py --spec-only T-12 --xlsx ... --token-file ...   # refill spec rows
    python nocodb_client.py --structure-only T-12 --token-file ...          # add missing columns + views
    python nocodb_client.py --compare-view-only T-12 --token-file ...       # (re)build "Price comparison"
    python nocodb_client.py --docs-list-only T-12 --docs-list-file docs.txt --token-file ...

Output (stdout, JSON):
    {"action": "created"|"updated", "record_id": 17, "purchase_number": "...",
     "verification": {"ok": true, "checked_fields": 38, "issues": []},
     "spec_table_url": "...", "spec_verification": {"ok": true, ...}}

"ok": false never changes the exit code (the write itself succeeded); the
caller must read "issues" and report them to the user.
"""

import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import load_config  # noqa: E402
from save_tender import (NOT_SPECIFIED, NUMBER_TOKEN_RE, apply_dash_placeholders,  # noqa: E402
                         format_source_files, parse_number_token)

CFG = load_config()
NC = CFG["nocodb"]
NOCODB_URL = NC["url"].rstrip("/")
API_URL = f"{NOCODB_URL}/api/v2"
TABLES_URL = f"{API_URL}/tables"
META_URL = f"{API_URL}/meta"
STORAGE_UPLOAD_URL = f"{API_URL}/storage/upload"
WORKSPACE_ID = NC.get("workspace_id", "")
BASE_ID = NC.get("base_id", "")
REGISTRY_TABLE_ID = NC.get("registry_table_id", "")
BASE_CURRENCY = str(CFG.get("base_currency") or "EUR").upper()
NOCODB_MAX_BATCH = 100  # bulk insert/update/delete cap per request

DECIMAL_META = {"dtxp": 15, "dtxs": 2, "meta": {"precision": 2}}
QTY_META = {"dtxp": 15, "dtxs": 3, "meta": {"precision": 3}}

# =============================================================================
# Registry table (one row per tender)
# =============================================================================
# flat field name (same keys as save_tender.py) -> (column title, NocoDB type)
FIELD_MAP = {
    "internal_number": ("Internal number", "SingleLineText"),
    "purchase_number": ("Tender number", "SingleLineText"),
    "tender_url": ("Tender URL", "URL"),
    "source_file": ("Source files", "LongText"),
    "extracted_date": ("Extracted on", "SingleLineText"),
    "product_name": ("Subject of procurement", "LongText"),
    "quantity": ("Quantity", "LongText"),
    "condition": ("Condition", "SingleLineText"),
    "delivery_place": ("Delivery place", "LongText"),
    "delivery_term": ("Delivery term", "LongText"),
    "procedure_type": ("Procedure type", "SingleLineText"),
    "price_reduction": ("Price-reduction stage", "SingleLineText"),
    "price_reduction_date": ("Price-reduction date", "SingleLineText"),
    "start_price": ("Starting price", "SingleLineText"),
    "currency": ("Currency", "SingleLineText"),
    "payment_terms": ("Payment terms", "LongText"),
    "bid_security": ("Bid security", "LongText"),
    "contract_security": ("Performance security", "LongText"),
    "publication_date": ("Publication date", "SingleLineText"),
    "submission_deadline": ("Submission deadline", "SingleLineText"),
    "auction_date": ("Auction date", "SingleLineText"),
    "results_date": ("Award date", "SingleLineText"),
    "contract_date": ("Contract signing", "LongText"),
    "last_modified": ("Last modified", "SingleLineText"),
    "buyer": ("Buyer", "SingleLineText"),
    "buyer_registration_id": ("Buyer registration ID", "SingleLineText"),
    "contact": ("Buyer contact", "LongText"),
    "address": ("Buyer address", "LongText"),
    "procurement_method": ("Procurement method", "SingleLineText"),
    "platform": ("Platform", "SingleLineText"),
    "platform_url": ("Platform URL", "URL"),
    "stage": ("Stage", "SingleLineText"),
    "sme": ("SME", "SingleLineText"),
    "domestic_preference_applied": ("Domestic preference: applied", "SingleLineText"),
    "domestic_preference_measure": ("Domestic preference: measure", "SingleLineText"),
    "domestic_preference_items": ("Domestic preference: items", "LongText"),
    "domestic_preference_basis": ("Domestic preference: basis", "LongText"),
    # NocoDB-only fields (not part of the XML/CSV):
    "payment_days_summary": ("Payment term (short)", "SingleLineText"),   # step 7.1a
    "application_documents": ("Required bid documents", "LongText"),       # step 7.1d
}
REGISTRY_COLUMN = {flat: title for flat, (title, _) in FIELD_MAP.items()}
INTERNAL_NUMBER_COLUMN = REGISTRY_COLUMN["internal_number"]
PURCHASE_NUMBER_COLUMN = REGISTRY_COLUMN["purchase_number"]
SOURCE_FILE_COLUMN = REGISTRY_COLUMN["source_file"]
APPLICATION_DOCS_COLUMN = REGISTRY_COLUMN["application_documents"]
SPEC_LINK_COLUMN = "Specification link"
SUPPLIER_FILES_COLUMN = "Supplier request files"
# Rich-text (Markdown) display for multi-line list fields.
RICH_TEXT_COLUMNS = {REGISTRY_COLUMN["product_name"], SOURCE_FILE_COLUMN, APPLICATION_DOCS_COLUMN}
# URL columns validate their value - placeholders must not be sent there.
STRICT_URL_FIELDS = ("tender_url", "platform_url")
_URL_RE = re.compile(r"^https?://", re.IGNORECASE)

# =============================================================================
# Specification table (one per tender) - columns of the "Items & prices" sheet
# =============================================================================
COL_NO = "No."
COL_ITEM = "Item name"
COL_TECH = "Technical requirements"
COL_DOC_ERRORS = "Documentation errors"
COL_MODEL = "Brand/model"
COL_MAKER = "Manufacturer"
COL_ORIGIN = "Country of origin"
COL_EQ_ALLOWED = "Equivalent allowed"
COL_EQ_STATUS = "Equivalent (status)"
COL_EQUIVALENT = "Equivalent"
COL_QTY = "Qty"
COL_UNIT = "Unit"
COL_LINE_MAX = "Max price (line)"
COL_PACKAGE = "Package dimensions"
COL_URL = "Manufacturer URL"
COL_UNIT_MAX = "Max unit price"
COL_NOTES = "Notes"

EQ_OPTIONS = [
    {"title": "Yes", "color": "#6bdc8b"},     # green - an equivalent may be offered
    {"title": "No", "color": "#ff6b6b"},      # red - equivalents are not allowed
    {"title": "Check", "color": "#ffd43b"},   # yellow - documents unclear / contradictory
]

# Agent prices (step 5.1f): up to three offers, each followed by its comment.
AGENT_PRICE_PAIRS = [(f"Agent price {i}", f"Agent comment {i}") for i in (1, 2, 3)]
AGENT_PRICE_COLUMNS = tuple(p for p, _ in AGENT_PRICE_PAIRS)
AGENT_FIELDS = [COL_UNIT_MAX] + [c for pair in AGENT_PRICE_PAIRS for c in pair]

# Market prices (step 5.1g): three confirmed offers per market, price then link.
MARKETS = [("dom", "Domestic"), ("intl", "International")]
MARKET_OFFER_PAIRS = [(f"{label} offer {i}", f"{label} link {i}") for _, label in MARKETS for i in (1, 2, 3)]
MARKET_OFFER_PRICE_COLUMNS = tuple(p for p, _ in MARKET_OFFER_PAIRS)
MARKET_OFFER_LINK_COLUMNS = tuple(lnk for _, lnk in MARKET_OFFER_PAIRS)
MARKET_OFFER_FIELDS = [c for pair in MARKET_OFFER_PAIRS for c in pair]

# Procurement-team prices (step 5.1h) - filled by people or the Telegram bot,
# never by the agent; preserved by "No." when the spec table is refilled.
PROCUREMENT_SLOTS = int(CFG.get("procurement_slots") or 10)
PROCUREMENT_PAIRS = [(f"Procurement: supplier {i}", f"Procurement: price {i}")
                     for i in range(1, PROCUREMENT_SLOTS + 1)]
PROCUREMENT_SUPPLIER_COLUMNS = tuple(s for s, _ in PROCUREMENT_PAIRS)
PROCUREMENT_PRICE_COLUMNS = tuple(p for _, p in PROCUREMENT_PAIRS)
PROCUREMENT_COMMENT_COLUMN = "Procurement: comment"
PROCUREMENT_FIELDS = [c for pair in PROCUREMENT_PAIRS for c in pair] + [PROCUREMENT_COMMENT_COLUMN]


def _db_name(title):
    """Stable snake_case physical column name derived from a title."""
    return re.sub(r"[^a-z0-9]+", "_", title.lower()).strip("_")[:60] or "col"


SPEC_COLUMNS = [
    (COL_NO, "SingleLineText", None),
    (COL_ITEM, "LongText", None),
    (COL_TECH, "LongText", None),
    (COL_DOC_ERRORS, "LongText", None),
    (COL_MODEL, "SingleLineText", None),
    (COL_MAKER, "SingleLineText", None),
    (COL_ORIGIN, "SingleLineText", None),
    (COL_EQ_ALLOWED, "SingleLineText", None),
    (COL_EQ_STATUS, "SingleSelect", {"colOptions": {"options": EQ_OPTIONS}}),
    (COL_EQUIVALENT, "LongText", None),
    (COL_QTY, "Decimal", QTY_META),
    (COL_UNIT, "SingleLineText", None),
    (COL_LINE_MAX, "Decimal", DECIMAL_META),
    (COL_PACKAGE, "LongText", None),
    (COL_URL, "URL", None),
    (COL_UNIT_MAX, "Decimal", DECIMAL_META),
] + [(c, "Decimal" if c in AGENT_PRICE_COLUMNS else "LongText",
      DECIMAL_META if c in AGENT_PRICE_COLUMNS else None) for pair in AGENT_PRICE_PAIRS for c in pair] \
  + [(c, "URL" if c in MARKET_OFFER_LINK_COLUMNS else "Decimal",
      None if c in MARKET_OFFER_LINK_COLUMNS else DECIMAL_META) for c in MARKET_OFFER_FIELDS] \
  + [(c, "Decimal" if c in PROCUREMENT_PRICE_COLUMNS else
      ("LongText" if c == PROCUREMENT_COMMENT_COLUMN else "SingleLineText"),
      DECIMAL_META if c in PROCUREMENT_PRICE_COLUMNS else None) for c in PROCUREMENT_FIELDS] \
  + [(COL_NOTES, "LongText", None)]

SPEC_TITLES = [name for name, _, _ in SPEC_COLUMNS]
NUMERIC_SPEC_COLUMNS = (COL_QTY, COL_LINE_MAX, COL_UNIT_MAX) + AGENT_PRICE_COLUMNS \
    + MARKET_OFFER_PRICE_COLUMNS + PROCUREMENT_PRICE_COLUMNS

# --- views --------------------------------------------------------------------
VIEW_LEAD_COLUMNS = [COL_NO, COL_ITEM, COL_MODEL, COL_QTY, COL_UNIT]
AGENT_VIEW_TITLE = "Agent prices"
PROCUREMENT_VIEW_TITLE = "Procurement"
PROCUREMENT_VIEW_FIELDS = PROCUREMENT_FIELDS + [COL_EQ_ALLOWED, COL_EQ_STATUS, COL_MAKER, COL_EQUIVALENT,
                                                COL_ORIGIN, COL_DOC_ERRORS] + MARKET_OFFER_FIELDS
COMPARE_VIEW_TITLE = "Price comparison"
COMPARE_BEST_PRICE_COLUMNS = [f"Procurement best price {k}" for k in (1, 2, 3)]
COMPARE_BEST_SUPPLIER_COLUMNS = [f"Procurement best supplier {k}" for k in (1, 2, 3)]
COMPARE_PURCH_TEXT_COLUMNS = [f"Procurement price {k}" for k in (1, 2, 3)]
COMPARE_MARKET_TEXT_COLUMNS = {key: [f"{label} price {i}" for i in (1, 2, 3)] for key, label in MARKETS}
COMPARE_VIEW_FIELDS = (
    [COL_UNIT_MAX]
    + [c for pair in zip(COMPARE_PURCH_TEXT_COLUMNS, COMPARE_BEST_SUPPLIER_COLUMNS) for c in pair]
    + [c for key, label in MARKETS for i in (1, 2, 3)
       for c in (COMPARE_MARKET_TEXT_COLUMNS[key][i - 1], f"{label} link {i}")]
)

# =============================================================================
# Value helpers
# =============================================================================
# "120.00 per unit (240.00 for 2 pcs)" -> the line total in parentheses.
POSITION_TOTAL_RE = re.compile(r"\(([\d\s.,']+)\s*(?:for|per)\s*\d+\s*(?:pcs|units?|ea|items?)\.?\)", re.IGNORECASE)

# Currency symbols / codes recognised in price cells.
CURRENCY_SYMBOLS = {"€": "EUR", "$": "USD", "£": "GBP", "¥": "CNY", "₽": "RUB", "₹": "INR", "₺": "TRY"}
CURRENCY_CODE_RE = re.compile(r"\b(EUR|USD|GBP|CNY|RMB|JPY|CHF|RUB|INR|TRY|PLN|CZK|SEK|NOK|DKK|RSD|HUF|AED|CAD|AUD)\b",
                              re.IGNORECASE)


def parse_number(raw):
    """First number in a cell (int/float or text in any common notation), or None."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    match = NUMBER_TOKEN_RE.search(str(raw))
    return parse_number_token(match.group(0)) if match else None


def parse_position_total(raw):
    """"Max price (line)" prefers an explicit "(X for N pcs)" total over a
    leading per-unit figure."""
    if isinstance(raw, (int, float)) or raw is None:
        return parse_number(raw)
    match = POSITION_TOTAL_RE.search(str(raw))
    return parse_number(match.group(1)) if match else parse_number(raw)


def detect_currency(text):
    """ISO code of the currency mentioned in a price cell, or None."""
    text = str(text or "")
    for sym, code in CURRENCY_SYMBOLS.items():
        if sym in text:
            return code
    m = CURRENCY_CODE_RE.search(text)
    if m:
        code = m.group(1).upper()
        return "CNY" if code == "RMB" else code
    return None


def convert_to_base(raw_text, rates=None):
    """("<value in base currency>" | None, note | None) for one price cell.

    rates: {"USD": 0.92, ...} = base-currency units per 1 unit of that
    currency (pass with --rate USD=0.92). A foreign price with no rate is
    left empty and explained in the note - never guessed.
    """
    if raw_text is None or str(raw_text).strip() == "":
        return None, None
    value = parse_number(raw_text)
    if value is None:
        return None, None
    currency = detect_currency(raw_text)
    if currency is None or currency == BASE_CURRENCY:
        return round(value, 2), None
    rate = (rates or {}).get(currency)
    if rate is None:
        return None, (f"not converted - no {currency}/{BASE_CURRENCY} rate given "
                      f"(--rate {currency}=...) for \"{raw_text}\"")
    return round(value * rate, 2), (f"converted from {value:.2f} {currency} at {rate} {BASE_CURRENCY}/{currency} "
                                    f"({date.today().isoformat()})")


def _append(existing, addition, sep="; "):
    existing = str(existing or "").strip()
    return f"{existing}{sep}{addition}" if existing else addition


def convert_foreign_prices(records, rates=None):
    """Agent prices and market offers written as "120 USD"/"$45.90" are
    converted into the base currency; the conversion is explained in the
    matching comment (agent prices) or in Notes (market offers)."""
    out = []
    for rec in records:
        rec = dict(rec)
        for price_col, comment_col in AGENT_PRICE_PAIRS:
            raw = rec.get(price_col)
            if isinstance(raw, str) and detect_currency(raw) not in (None, BASE_CURRENCY):
                rec[price_col], note = convert_to_base(raw, rates)
                if note:
                    rec[comment_col] = _append(rec.get(comment_col), f"{price_col}: {note}", "\n")
        for price_col in MARKET_OFFER_PRICE_COLUMNS:
            raw = rec.get(price_col)
            if isinstance(raw, str) and detect_currency(raw) not in (None, BASE_CURRENCY):
                rec[price_col], note = convert_to_base(raw, rates)
                if note:
                    rec[COL_NOTES] = _append(rec.get(COL_NOTES), f"{price_col}: {note}")
        out.append(rec)
    return out


def normalize_equivalent_status(raw):
    """Maps an "Equivalent (status)"/"Equivalent allowed" value onto Yes/No/Check.
    Anything that isn't clearly Yes or No becomes "Check" - never guess."""
    low = str(raw or "").strip().lower()
    if low == "yes" or re.match(r"^yes\b", low):
        return "Yes"
    if low == "no" or re.match(r"^no\b", low):
        return "No"
    return "Check"


def sanitize_spec_row(row):
    """One xlsx row -> a record NocoDB accepts.

    Numeric columns must be numbers or null: a placeholder in any of them
    makes NocoDB reject the WHOLE bulk insert. Unparseable text is moved to
    Notes instead of being lost. Also flags prices without comment/link/supplier.
    """
    out = dict(row)
    notes = []
    status = out.get(COL_EQ_STATUS)
    if str(status or "").strip() not in {o["title"] for o in EQ_OPTIONS}:
        status = out.get(COL_EQ_ALLOWED)
    out[COL_EQ_STATUS] = normalize_equivalent_status(status)
    for col in NUMERIC_SPEC_COLUMNS:
        if col not in out:
            continue
        raw = out[col]
        value = parse_position_total(raw) if col == COL_LINE_MAX else parse_number(raw)
        if value is None and raw is not None and str(raw).strip():
            notes.append(f"{col}: {raw}")
        out[col] = value
    for price_col, comment_col in AGENT_PRICE_PAIRS:
        if out.get(price_col) is not None and not str(out.get(comment_col) or "").strip():
            notes.append(f"⚠ {price_col} has no comment (source/link missing)")
    for price_col, link_col in MARKET_OFFER_PAIRS:
        if out.get(price_col) is not None and not str(out.get(link_col) or "").strip():
            notes.append(f"⚠ {price_col} has no link")
    for supplier_col, price_col in PROCUREMENT_PAIRS:
        if out.get(price_col) is not None and not str(out.get(supplier_col) or "").strip():
            notes.append(f"⚠ {price_col} has no supplier")
    if notes:
        out[COL_NOTES] = _append(out.get(COL_NOTES), "; ".join(notes))
    return out


def read_spec_xlsx(xlsx_path):
    """First sheet of Items_and_prices.xlsx -> one dict per real item row
    (rows without an item name - totals, blank or note rows - are skipped)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    if not rows:
        return []
    headers = rows[0]
    records = []
    for r in rows[1:]:
        rec = {h: v for h, v in zip(headers, r) if v is not None and h is not None}
        if str(rec.get(COL_ITEM) or "").strip():
            records.append(rec)
    return records


# =============================================================================
# HTTP / NocoDB primitives
# =============================================================================
def _request(method, url, token, body=None, retries=8):
    """JSON request with retry/back-off on HTTP 429 (NocoDB allows ~5 req/s)."""
    data = json.dumps(body).encode("utf-8") if body is not None else None
    for attempt in range(retries + 1):
        if url.startswith(META_URL):
            time.sleep(0.3)
        req = urllib.request.Request(url, data=data, method=method)
        req.add_header("xc-token", token)
        req.add_header("Content-Type", "application/json")
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                raw = resp.read().decode("utf-8")
                return resp.status, (json.loads(raw) if raw else {})
        except urllib.error.HTTPError as e:
            raw = e.read().decode("utf-8", errors="replace")
            if e.code == 429 and attempt < retries:
                time.sleep(min(60, 5 * (attempt + 1)))
                continue
            return e.code, {"error": raw}
    return 429, {"error": "too many requests"}


def _ok(status, body, what):
    if status != 200:
        raise RuntimeError(f"{what} failed: HTTP {status} {body}")
    return body


def _records_url(table_id):
    return f"{TABLES_URL}/{table_id}/records"


def _require_registry():
    if not REGISTRY_TABLE_ID:
        raise RuntimeError("nocodb.registry_table_id is not set in config.json - run --init-registry first")


def find_record(token, column, value, table_id=None):
    table_id = table_id or REGISTRY_TABLE_ID
    where = urllib.parse.quote(f"({column},eq,{value})")
    status, body = _request("GET", f"{_records_url(table_id)}?where={where}&limit=1", token)
    rows = _ok(status, body, "lookup").get("list", [])
    return rows[0] if rows else None


def find_record_any(token, key):
    """Registry record by tender number, then by internal number."""
    return find_record(token, PURCHASE_NUMBER_COLUMN, key) or find_record(token, INTERNAL_NUMBER_COLUMN, key)


def get_table_columns(token, table_id):
    """{column title: column meta} of a table."""
    status, body = _request("GET", f"{META_URL}/tables/{table_id}", token)
    return {c["title"]: c for c in _ok(status, body, "get table").get("columns", [])}


def list_tables(token):
    status, body = _request("GET", f"{META_URL}/bases/{BASE_ID}/tables", token)
    return _ok(status, body, "list tables").get("list", [])


def find_table_by_title(token, title):
    return next((t for t in list_tables(token) if t.get("title") == title), None)


def list_views(token, table_id):
    status, body = _request("GET", f"{META_URL}/tables/{table_id}/views", token)
    return _ok(status, body, "list views").get("list", [])


def get_default_view_id(token, table_id):
    views = list_views(token, table_id)
    return views[0]["id"] if views else None


def list_all_records(token, table_id):
    """All rows of a table (NocoDB returns at most 100 per page)."""
    rows, offset = [], 0
    while True:
        status, body = _request("GET", f"{_records_url(table_id)}?limit=100&offset={offset}", token)
        batch = _ok(status, body, "list records").get("list", [])
        rows.extend(batch)
        if not batch or body.get("pageInfo", {}).get("isLastPage", True):
            return rows
        offset += len(batch)


def position_column_after(token, table_id, column_id, anchor_title):
    """Best-effort: moves a column right after `anchor_title` in the default
    grid view. Column order is stored per view, on the view-column record
    (not on the column itself); orders are floats, so the midpoint works."""
    try:
        view_id = get_default_view_id(token, table_id)
        status, body = _request("GET", f"{META_URL}/views/{view_id}/columns", token)
        if status != 200:
            return
        view_columns = body.get("list", [])
        anchor_id = get_table_columns(token, table_id).get(anchor_title, {}).get("id")
        by_column = {vc["fk_column_id"]: vc for vc in view_columns}
        target, anchor = by_column.get(column_id), by_column.get(anchor_id)
        if not target or not anchor:
            return
        later = [vc["order"] for vc in view_columns if vc["order"] > anchor["order"]]
        new_order = (anchor["order"] + min(later)) / 2 if later else anchor["order"] + 1
        _request("PATCH", f"{META_URL}/views/{view_id}/columns/{target['id']}", token, {"order": new_order})
    except Exception:
        pass


def _column_body(title, uidt, extra=None):
    body = {"title": title, "uidt": uidt, "column_name": _db_name(title)}
    if extra:
        body.update(extra)
    if title in RICH_TEXT_COLUMNS and uidt == "LongText":
        body["meta"] = {"richMode": True}
    return body


# =============================================================================
# Registry
# =============================================================================
def registry_column_defs():
    defs = [(title, uidt) for title, uidt in FIELD_MAP.values()]
    defs += [(SPEC_LINK_COLUMN, "URL"), (SUPPLIER_FILES_COLUMN, "Attachment")]
    return defs


def init_registry(token):
    """Creates the registry table (if registry_table_id is empty) or adds any
    missing registry columns to the existing one. Idempotent."""
    if not BASE_ID:
        raise RuntimeError("nocodb.base_id is not set in config.json")
    title = NC.get("registry_table_title") or "Tender Registry"
    table_id = REGISTRY_TABLE_ID
    created = []
    if not table_id:
        existing = find_table_by_title(token, title)
        if existing:
            table_id = existing["id"]
        else:
            cols = [_column_body(t, u) for t, u in registry_column_defs() if u != "Attachment"]
            cols[0]["pv"] = True  # "Internal number" is the display (primary) column
            body = _ok(*_request("POST", f"{META_URL}/bases/{BASE_ID}/tables", token,
                                 {"title": title, "columns": cols}), "create registry table")
            table_id = body["id"]
            created = [c["title"] for c in cols]
    have = get_table_columns(token, table_id)
    prev = None
    for col_title, uidt in registry_column_defs():
        if col_title not in have:
            _ok(*_request("POST", f"{META_URL}/tables/{table_id}/columns", token,
                          _column_body(col_title, uidt)), f"create column '{col_title}'")
            created.append(col_title)
            if prev:
                new_id = get_table_columns(token, table_id).get(col_title, {}).get("id")
                if new_id:
                    position_column_after(token, table_id, new_id, prev)
        prev = col_title
    return {"action": "init_registry", "registry_table_id": table_id, "created_columns": created,
            "hint": None if REGISTRY_TABLE_ID else f'set "registry_table_id": "{table_id}" in config.json'}


def ensure_registry_columns(token, titles):
    """Auto-creates registry columns used by this run that are missing."""
    wanted = {title: uidt for title, uidt in registry_column_defs() if title in titles}
    have = get_table_columns(token, REGISTRY_TABLE_ID)
    for title, uidt in wanted.items():
        if title not in have:
            _ok(*_request("POST", f"{META_URL}/tables/{REGISTRY_TABLE_ID}/columns", token,
                          _column_body(title, uidt)), f"create column '{title}'")


def _looks_like_url(value):
    return bool(_URL_RE.match(str(value or "").strip()))


def build_registry_row(flat_data):
    apply_dash_placeholders(flat_data)
    if flat_data.get("source_file"):
        flat_data["source_file"] = format_source_files(str(flat_data["source_file"]))
    row = {}
    for flat_name, value in flat_data.items():
        if flat_name not in REGISTRY_COLUMN:
            continue  # unknown key - skip rather than fail
        if flat_name in STRICT_URL_FIELDS and not _looks_like_url(value):
            continue  # "—"/placeholder would fail URL validation for the whole record
        row[REGISTRY_COLUMN[flat_name]] = value
    return row


def upsert(token, flat_data, spec_xlsx_path=None, fallback_id=None, rates=None):
    """Creates/updates the registry record; with spec_xlsx_path also syncs the
    tender's specification table. fallback_id names the spec table when the
    tender has neither an internal nor an official tender number."""
    _require_registry()
    purchase_number = flat_data.get("purchase_number")
    if not purchase_number:
        raise ValueError("the JSON must include 'purchase_number' (or the NOT_SPECIFIED placeholder)")
    row = build_registry_row(flat_data)
    ensure_registry_columns(token, set(row) | ({SPEC_LINK_COLUMN} if spec_xlsx_path else set()))

    # Without a real tender number the placeholder would match unrelated
    # tenders - match by source files instead.
    if purchase_number == NOT_SPECIFIED:
        source_file = row.get(SOURCE_FILE_COLUMN)
        existing = find_record(token, SOURCE_FILE_COLUMN, source_file) if source_file else None
    else:
        existing = find_record(token, PURCHASE_NUMBER_COLUMN, purchase_number)

    spec_table_url = spec_table_id = spec_rows_sent = None
    if spec_xlsx_path:
        internal = str(flat_data.get("internal_number") or "").strip()
        if internal and internal != NOT_SPECIFIED:
            spec_id = internal
        elif purchase_number != NOT_SPECIFIED:
            spec_id = purchase_number
        else:
            spec_id = fallback_id
        if spec_id:
            rows = convert_foreign_prices(read_spec_xlsx(spec_xlsx_path), rates)
            linked = table_id_from_link((existing or {}).get(SPEC_LINK_COLUMN))
            spec_table_url, spec_table_id, spec_rows_sent = sync_specification(token, spec_id, rows, linked)
            row[SPEC_LINK_COLUMN] = spec_table_url

    if existing:
        record_id = existing["Id"]
        _ok(*_request("PATCH", _records_url(REGISTRY_TABLE_ID), token, {"Id": record_id, **row}), "update")
        result = {"action": "updated", "record_id": record_id, "purchase_number": purchase_number}
    else:
        body = _ok(*_request("POST", _records_url(REGISTRY_TABLE_ID), token, row), "create")
        result = {"action": "created", "record_id": body.get("Id"), "purchase_number": purchase_number}

    # Trust nothing: re-read what was written and compare.
    result["verification"] = verify_main_record(token, result["record_id"], row)
    if spec_table_url:
        result["spec_table_url"] = spec_table_url
        result["spec_verification"] = verify_spec_table(token, spec_table_id, spec_rows_sent or [])
    return result


# =============================================================================
# Specification tables
# =============================================================================
def _safe_title(value):
    # NocoDB rejects "/" and "." in table titles.
    return str(value).strip().replace("/", "-").replace(".", "-")


def table_id_from_link(url):
    """Table id from a "Specification link" URL (.../<base_id>/<table_id>[/<view_id>])."""
    if not url or not BASE_ID:
        return None
    m = re.search(re.escape(BASE_ID) + r"/(m[a-z0-9]+)", str(url))
    return m.group(1) if m else None


def table_url(table_id, view_id=None):
    url = f"{NOCODB_URL}/{WORKSPACE_ID}/{BASE_ID}/{table_id}"
    return f"{url}/{view_id}" if view_id else url


def rename_table(token, table_id, title):
    status, meta = _request("GET", f"{META_URL}/tables/{table_id}", token)
    meta = _ok(status, meta, "get table")
    if meta.get("title") != title:
        _ok(*_request("PATCH", f"{META_URL}/tables/{table_id}", token,
                      {"title": title, "table_name": meta["table_name"]}), f"rename table to '{title}'")
        meta["title"] = title
    return meta


def create_spec_table(token, title):
    columns = [_column_body(name, uidt, extra) for name, uidt, extra in SPEC_COLUMNS]
    return _ok(*_request("POST", f"{META_URL}/bases/{BASE_ID}/tables", token,
                         {"title": title, "columns": columns}), "create spec table")


def ensure_spec_columns(token, table_id, only=None):
    """Adds SPEC_COLUMNS missing from an existing spec table (each placed after
    its predecessor). `only` restricts this to a subset of titles."""
    existing = get_table_columns(token, table_id)
    prev = None
    for name, uidt, extra in SPEC_COLUMNS:
        if name not in existing and (only is None or name in only):
            _ok(*_request("POST", f"{META_URL}/tables/{table_id}/columns", token,
                          _column_body(name, uidt, extra)), f"create spec column '{name}'")
            if prev:
                new_id = get_table_columns(token, table_id).get(name, {}).get("id")
                if new_id:
                    position_column_after(token, table_id, new_id, prev)
        prev = name


def clear_table_records(token, table_id):
    ids = [{"Id": r["Id"]} for r in list_all_records(token, table_id)]
    for i in range(0, len(ids), NOCODB_MAX_BATCH):
        _ok(*_request("DELETE", _records_url(table_id), token, ids[i:i + NOCODB_MAX_BATCH]), "clear records")


def read_manual_fields(token, table_id):
    """{No.: {procurement field: value}} - values entered by people/the bot,
    saved before the table is refilled from xlsx."""
    saved = {}
    for r in list_all_records(token, table_id):
        values = {c: r.get(c) for c in PROCUREMENT_FIELDS if r.get(c) not in (None, "")}
        key = str(r.get(COL_NO) or "").strip()
        if values and key:
            saved[key] = values
    return saved


def restore_manual_fields(rows, manual):
    """Puts saved procurement values back by "No." (xlsx values win if present)."""
    out = []
    for r in rows:
        r = dict(r)
        for col, val in manual.get(str(r.get(COL_NO) or "").strip(), {}).items():
            if r.get(col) in (None, ""):
                r[col] = val
        out.append(r)
    return out


def ensure_view(token, table_id, view_title, fields):
    """Grid view showing VIEW_LEAD_COLUMNS + `fields` in this order, everything
    else hidden. Created if missing, re-applied on every run (idempotent).
    Returns {"ok", "view_id", "issues"} after re-reading the view."""
    wanted = VIEW_LEAD_COLUMNS + list(fields)
    try:
        view = next((v for v in list_views(token, table_id) if v.get("title") == view_title), None)
        if view is None:
            view = _ok(*_request("POST", f"{META_URL}/tables/{table_id}/grids", token, {"title": view_title}),
                       "create view")
        view_id = view["id"]
        id_to_title = {meta["id"]: title for title, meta in get_table_columns(token, table_id).items()}
        body = _ok(*_request("GET", f"{META_URL}/views/{view_id}/columns", token), "list view columns")
        for vc in body.get("list", []):
            title = id_to_title.get(vc["fk_column_id"])
            if title in wanted:
                patch = {"show": True, "order": wanted.index(title) + 1}
            else:
                patch = {"show": False, "order": len(wanted) + 1 + (vc.get("order") or 0)}
            if bool(vc.get("show")) != patch["show"] or vc.get("order") != patch["order"]:
                _ok(*_request("PATCH", f"{META_URL}/views/{view_id}/columns/{vc['id']}", token, patch),
                    f"update view column '{title}'")
        body = _ok(*_request("GET", f"{META_URL}/views/{view_id}/columns", token), "list view columns")
        shown = [id_to_title.get(vc["fk_column_id"]) for vc in
                 sorted(body.get("list", []), key=lambda vc: vc.get("order") or 0) if vc.get("show")]
        issues = [] if shown == wanted else [f"visible fields: {shown}, expected: {wanted}"]
        return {"ok": not issues, "view_id": view_id, "issues": issues}
    except Exception as e:
        return {"ok": False, "view_id": None, "issues": [str(e)]}


def ensure_spec_views(token, table_id):
    results = {
        AGENT_VIEW_TITLE: ensure_view(token, table_id, AGENT_VIEW_TITLE, AGENT_FIELDS),
        PROCUREMENT_VIEW_TITLE: ensure_view(token, table_id, PROCUREMENT_VIEW_TITLE, PROCUREMENT_VIEW_FIELDS),
    }
    for title, res in results.items():
        if not res.get("ok"):
            print(f"WARNING: view '{title}': {res.get('issues')}", file=sys.stderr)
    return results


# --- "Price comparison": formula columns -------------------------------------------
def _ref(title):
    return "{" + title + "}"


def compare_formula_columns():
    """[(title, formula, meta, column_name)] for the "Price comparison" view.

    NocoDB has no SMALL()/LARGE(), so the k-th best procurement price is found
    through each price's rank: rank = (# lower prices) + (# EQUAL prices in
    lower-numbered slots), so ties are never lost. Empty prices are ignored.
    The text columns render "38,164.00 🏆" - the trophy marks the lowest price
    of each group (procurement / domestic / international), all ties included.
    """
    prices = [_ref(p) for _, p in PROCUREMENT_PAIRS]
    suppliers = [_ref(s) for s, _ in PROCUREMENT_PAIRS]
    n = len(prices)

    def rank(i):
        terms = [f"IF({prices[j]} < {prices[i]}, 1, 0)" for j in range(n) if j != i]
        terms += [f"IF({prices[j]} == {prices[i]}, 1, 0)" for j in range(i)]
        return "(" + " + ".join(terms) + ")" if terms else "0"

    def kth(k, values, fallback):
        expr = fallback
        for i in reversed(range(n)):
            expr = f"IF(AND({prices[i]} > 0, {rank(i)} == {k}), {values[i]}, {expr})"
        return expr

    def money_text(x):
        return f'REGEX_REPLACE(CONCAT(ROUND({x}, 2), ""), "(\\d)(?=(\\d{{3}})+(\\.|$))", "\\1,")'

    def marked(x, group_min, searched=None):
        # searched: "this market was searched" -> an empty price reads "not found"
        empty = f'IF({searched}, "not found", "")' if searched else '""'
        return f'IF({x} > 0, CONCAT({money_text(x)}, IF({x} == {group_min}, " 🏆", "")), {empty})'

    money_meta = {"display_type": "Decimal", "display_column_meta": {"meta": dict(DECIMAL_META["meta"]), "custom": {}}}
    cols = []
    for k in range(3):
        cols.append((COMPARE_BEST_PRICE_COLUMNS[k], kth(k, prices, "BLANK()"), money_meta, f"cmp_purch_best_{k + 1}"))
        cols.append((COMPARE_BEST_SUPPLIER_COLUMNS[k], kth(k, suppliers, '""'), None, f"cmp_purch_supplier_{k + 1}"))
    best1 = _ref(COMPARE_BEST_PRICE_COLUMNS[0])
    for k in range(3):
        cols.append((COMPARE_PURCH_TEXT_COLUMNS[k], marked(_ref(COMPARE_BEST_PRICE_COLUMNS[k]), best1), None,
                     f"cmp_purch_text_{k + 1}"))
    for key, label in MARKETS:
        offers = [_ref(f"{label} offer {i}") for i in (1, 2, 3)]
        group_min = "MIN(" + ", ".join(offers) + ")"
        # "searched" = any price present OR the step 5.1g note "<Market>: found N of 3 [...]" /
        # "<Market>: not found [...]" in Notes (REGEX_MATCH returns 1/0).
        searched = ("OR(" + ", ".join(f"{o} > 0" for o in offers)
                    + f', REGEX_MATCH({_ref(COL_NOTES)}, "{label}: (found|not found)") == 1)')
        for i in (1, 2, 3):
            cols.append((COMPARE_MARKET_TEXT_COLUMNS[key][i - 1], marked(offers[i - 1], group_min, searched), None,
                         f"cmp_{key}_text_{i}"))
    return cols


def ensure_compare_table(token, table_id):
    """Adds missing procurement/market columns, creates or updates the formula
    columns and the "Price comparison" view. Rows are not touched."""
    before = set(get_table_columns(token, table_id))
    ensure_spec_columns(token, table_id, only=PROCUREMENT_FIELDS + MARKET_OFFER_FIELDS)
    columns = get_table_columns(token, table_id)
    issues = [f"missing column: {c}" for c in PROCUREMENT_FIELDS + MARKET_OFFER_FIELDS if c not in columns]
    formulas = []
    if not issues:
        for title, formula, meta, column_name in compare_formula_columns():
            body = {"title": title, "uidt": "Formula", "formula_raw": formula, "column_name": column_name}
            if meta:
                body["meta"] = meta
            existing = columns.get(title)
            if existing is None:
                status, resp = _request("POST", f"{META_URL}/tables/{table_id}/columns", token, body)
                action = "created"
            elif (existing.get("colOptions") or {}).get("formula_raw") != formula:
                status, resp = _request("PATCH", f"{META_URL}/columns/{existing['id']}", token, body)
                action = "updated"
            else:
                status, resp, action = 200, {}, "unchanged"
            if status != 200:
                issues.append(f"formula '{title}': HTTP {status} {resp}")
            formulas.append({"title": title, "action": action if status == 200 else "error"})
            columns = get_table_columns(token, table_id)
    view = ensure_view(token, table_id, COMPARE_VIEW_TITLE, COMPARE_VIEW_FIELDS) if not issues else \
        {"ok": False, "issues": ["view not configured - see issues above"]}
    if not view.get("ok"):
        issues += view.get("issues", [])
    created = [c for c in get_table_columns(token, table_id) if c not in before]
    return {"action": "compare_view", "table_id": table_id, "created_columns": created,
            "formulas": formulas, "view": view, "issues": issues, "ok": not issues}


def sync_specification(token, spec_id, spec_rows, existing_table_id=None):
    """Creates or refills the tender's spec table. Procurement values entered
    by people are preserved by "No.". Returns (url, table_id, rows_sent)."""
    title = _safe_title(spec_id)
    table, manual, created_now = None, {}, False
    if existing_table_id:
        try:
            table = rename_table(token, existing_table_id, title)
        except RuntimeError:
            table = None
    if table is None:
        table = find_table_by_title(token, title)
    if table is None:
        table = create_spec_table(token, title)
        created_now = True
    else:
        ensure_spec_columns(token, table["id"])
        manual = read_manual_fields(token, table["id"])
        clear_table_records(token, table["id"])

    table_id = table["id"]
    rows_sent = [sanitize_spec_row(r) for r in restore_manual_fields(spec_rows or [], manual)]
    for i in range(0, len(rows_sent), NOCODB_MAX_BATCH):
        _ok(*_request("POST", _records_url(table_id), token, rows_sent[i:i + NOCODB_MAX_BATCH]), "insert spec rows")

    ensure_spec_views(token, table_id)
    has_compare = any(v.get("title") == COMPARE_VIEW_TITLE for v in list_views(token, table_id))
    if created_now or has_compare:
        try:
            res = ensure_compare_table(token, table_id)
            if not res["ok"]:
                print(f"WARNING: view '{COMPARE_VIEW_TITLE}': {res['issues']}", file=sys.stderr)
        except Exception as e:
            print(f"WARNING: view '{COMPARE_VIEW_TITLE}': {e}", file=sys.stderr)
    return table_url(table_id, get_default_view_id(token, table_id)), table_id, rows_sent


def find_spec_table_id(token, spec_id):
    rec = find_record_any(token, spec_id)
    table_id = table_id_from_link(rec.get(SPEC_LINK_COLUMN)) if rec else None
    if table_id is None:
        t = find_table_by_title(token, _safe_title(spec_id))
        table_id = t["id"] if t else None
    if table_id is None:
        raise RuntimeError(f"specification table not found: {spec_id}")
    return table_id


def sync_spec_only(token, spec_id, xlsx_path, rates=None):
    """Refills only the spec table of an already saved tender."""
    rec = find_record_any(token, spec_id)
    title_id, linked = spec_id, None
    if rec:
        title_id = str(rec.get(INTERNAL_NUMBER_COLUMN) or "").strip() or spec_id
        linked = table_id_from_link(rec.get(SPEC_LINK_COLUMN))
    rows = convert_foreign_prices(read_spec_xlsx(xlsx_path), rates)
    url, table_id, rows_sent = sync_specification(token, title_id, rows, linked)
    if rec and rec.get(SPEC_LINK_COLUMN) != url:
        _request("PATCH", _records_url(REGISTRY_TABLE_ID), token, {"Id": rec["Id"], SPEC_LINK_COLUMN: url})
    return {"action": "spec_only", "spec_id": title_id, "spec_table_url": url,
            "spec_verification": verify_spec_table(token, table_id, rows_sent)}


def ensure_spec_structure(token, spec_id):
    """Adds missing spec columns and (re)builds the views; rows untouched."""
    table_id = find_spec_table_id(token, spec_id)
    before = set(get_table_columns(token, table_id))
    ensure_spec_columns(token, table_id)
    after = get_table_columns(token, table_id)
    missing = [name for name in SPEC_TITLES if name not in after]
    views = ensure_spec_views(token, table_id)
    return {"action": "structure_only", "table_id": table_id,
            "created_columns": [c for c in after if c not in before], "missing_columns": missing,
            "views": views, "ok": not missing and all(v.get("ok") for v in views.values())}


# =============================================================================
# Verification
# =============================================================================
def _norm(value):
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _values_match(expected, actual, numeric=False):
    if numeric:
        def to_float(v):
            try:
                return float(v) if v not in (None, "") else None
            except (TypeError, ValueError):
                return None
        e, a = to_float(expected), to_float(actual)
        if e is None or a is None:
            return e is None and a is None
        return abs(e - a) < 0.005
    return _norm(expected) == _norm(actual)


def _issue(problem, **kw):
    return {**kw, "problem": problem}


def verify_main_record(token, record_id, expected_row):
    """Re-reads the registry record and compares every column sent."""
    try:
        status, actual = _request("GET", f"{_records_url(REGISTRY_TABLE_ID)}/{record_id}", token)
    except Exception as e:
        status, actual = None, str(e)
    if status != 200:
        return {"ok": False, "checked_fields": 0,
                "issues": [_issue(f"could not re-read the record: {status} {actual}")]}
    issues = []
    for col, expected in expected_row.items():
        got = actual.get(col)
        if not _values_match(expected, got):
            issues.append(_issue("empty in NocoDB" if _norm(got) == "" else "value differs",
                                 column=col, expected=expected, actual=got))
    return {"ok": not issues, "checked_fields": len(expected_row), "issues": issues}


def verify_spec_table(token, table_id, expected_rows):
    """Re-reads the spec table: row count, then text and numeric columns row by row."""
    try:
        actual_rows = sorted(list_all_records(token, table_id), key=lambda r: r.get("Id", 0))
    except Exception as e:
        return {"ok": False, "expected_rows": len(expected_rows), "actual_rows": None,
                "issues": [_issue(f"could not re-read the specification table: {e}")]}
    issues = []
    if len(actual_rows) != len(expected_rows):
        issues.append(_issue(f"row count differs: sent {len(expected_rows)}, found {len(actual_rows)}"))
    else:
        text_columns = (COL_ITEM, COL_TECH, COL_DOC_ERRORS, COL_EQUIVALENT, COL_MAKER, COL_ORIGIN,
                        COL_EQ_ALLOWED, COL_EQ_STATUS, COL_URL) + tuple(c for _, c in AGENT_PRICE_PAIRS) \
            + MARKET_OFFER_LINK_COLUMNS + PROCUREMENT_SUPPLIER_COLUMNS + (PROCUREMENT_COMMENT_COLUMN,)
        for idx, (expected, actual) in enumerate(zip(expected_rows, actual_rows), start=1):
            for col in text_columns + NUMERIC_SPEC_COLUMNS:
                if not _values_match(expected.get(col), actual.get(col), numeric=col in NUMERIC_SPEC_COLUMNS):
                    issues.append(_issue("empty in NocoDB" if _norm(actual.get(col)) == "" else "value differs",
                                         row=idx, item=expected.get(COL_ITEM), column=col,
                                         expected=expected.get(col), actual=actual.get(col)))
    return {"ok": not issues, "expected_rows": len(expected_rows), "actual_rows": len(actual_rows),
            "issues": issues}


# =============================================================================
# Attachments and the required-documents list
# =============================================================================
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def upload_file(token, path):
    """Uploads one file to NocoDB storage; returns the attachment object(s)."""
    boundary = "----tde" + uuid.uuid4().hex
    filename = os.path.basename(path)
    mime = DOCX_MIME if filename.lower().endswith(".docx") else "application/octet-stream"
    with open(path, "rb") as f:
        payload = f.read()
    quoted = urllib.parse.quote(filename)
    head = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
            f"filename=\"{quoted}\"; filename*=UTF-8''{quoted}\r\nContent-Type: {mime}\r\n\r\n").encode("utf-8")
    req = urllib.request.Request(STORAGE_UPLOAD_URL, data=head + payload + f"\r\n--{boundary}--\r\n".encode(),
                                 method="POST")
    req.add_header("xc-token", token)
    req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            result = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"upload '{filename}' failed: HTTP {e.code} {e.read().decode('utf-8', 'replace')}")
    items = result if isinstance(result, list) else [result]
    for it in items:
        it["title"] = filename
        it.setdefault("mimetype", mime)
    return items


def attach_files(token, record_id, paths, column=SUPPLIER_FILES_COLUMN):
    """REPLACES the attachment cell with `paths` and verifies by re-reading."""
    ensure_registry_columns(token, {column})
    attachments = [a for p in paths for a in upload_file(token, p)]
    _ok(*_request("PATCH", _records_url(REGISTRY_TABLE_ID), token, {"Id": record_id, column: attachments}),
        "attach files")
    status, rec = _request("GET", f"{_records_url(REGISTRY_TABLE_ID)}/{record_id}", token)
    got = (rec.get(column) or []) if status == 200 else []
    if isinstance(got, str):
        try:
            got = json.loads(got)
        except ValueError:
            got = []
    got_titles = sorted(a.get("title", "") for a in got)
    want_titles = sorted(a["title"] for a in attachments)
    issues = [] if got_titles == want_titles else [
        _issue("files in NocoDB differ from the files sent", expected=want_titles, actual=got_titles)]
    return {"ok": not issues, "column": column, "files": want_titles, "issues": issues}


def set_application_documents(token, record_id, text):
    """Step 7.1d: writes the required-documents list and verifies it."""
    ensure_registry_columns(token, {APPLICATION_DOCS_COLUMN})
    row = {APPLICATION_DOCS_COLUMN: text}
    _ok(*_request("PATCH", _records_url(REGISTRY_TABLE_ID), token, {"Id": record_id, **row}), "update")
    return verify_main_record(token, record_id, row)


def _docx_in(folder):
    return sorted(os.path.join(folder, n) for n in os.listdir(folder)
                  if n.lower().endswith(".docx") and not n.startswith("~$"))


# =============================================================================
# CLI
# =============================================================================
def parse_rates(values):
    rates = {}
    for item in values or []:
        code, _, rate = item.partition("=")
        if not rate:
            raise SystemExit(f"--rate expects CODE=RATE, got '{item}'")
        rates[code.strip().upper()] = float(rate.replace(",", "."))
    return rates


def _print(result, warn_keys=()):
    for key in warn_keys:
        part = result.get(key) if key else result
        if isinstance(part, dict) and part.get("ok") is False:
            print(f"WARNING: verification problems - see \"{key or 'issues'}\" in the JSON output.", file=sys.stderr)
    print(json.dumps(result, ensure_ascii=False))


def main():
    p = argparse.ArgumentParser(description="Sync an extracted tender into NocoDB",
                                formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    p.add_argument("--token-file", required=True, help="File containing the NocoDB API token")
    p.add_argument("--json", help="Flat tender JSON (same file as for save_tender.py)")
    p.add_argument("--xlsx", help="Items_and_prices.xlsx - also sync the specification table")
    p.add_argument("--fallback-id", help="Spec table name when there is no internal or tender number")
    p.add_argument("--rate", action="append", default=[], metavar="CODE=RATE",
                   help=f"Exchange rate: {BASE_CURRENCY} per 1 unit of CODE, e.g. USD=0.92 (repeatable)")
    p.add_argument("--supplier-files-dir", help="Folder of per-manufacturer .docx files (step 7.1b)")
    p.add_argument("--init-registry", action="store_true", help="Create/complete the registry table and exit")
    p.add_argument("--spec-only", metavar="ID", help="Refill only the spec table of a saved tender (needs --xlsx)")
    p.add_argument("--structure-only", metavar="ID", help="Add missing spec columns + views; rows untouched")
    p.add_argument("--compare-view-only", metavar="ID", help="(Re)build the 'Price comparison' view; rows untouched")
    p.add_argument("--docs-list-only", metavar="ID", help="Write --docs-list-file into 'Required bid documents'")
    p.add_argument("--docs-list-file", help="UTF-8 text file with the required-documents list (step 7.1d)")
    args = p.parse_args()

    with open(args.token_file, "r", encoding="utf-8") as f:
        token = f.read().strip()
    rates = parse_rates(args.rate)

    if args.init_registry:
        _print(init_registry(token))
        return
    _require_registry()
    if args.structure_only:
        _print(ensure_spec_structure(token, args.structure_only), [None])
        return
    if args.compare_view_only:
        _print(ensure_compare_table(token, find_spec_table_id(token, args.compare_view_only)), [None])
        return
    if args.docs_list_only:
        if not args.docs_list_file:
            p.error("--docs-list-only requires --docs-list-file")
        with open(args.docs_list_file, "r", encoding="utf-8") as f:
            text = f.read().strip()
        rec = find_record_any(token, args.docs_list_only)
        if not rec:
            print(json.dumps({"error": f"record not found: {args.docs_list_only}"}))
            sys.exit(1)
        _print({"action": "docs_list_only", "record_id": rec["Id"],
                "verification": set_application_documents(token, rec["Id"], text)}, ["verification"])
        return
    if args.spec_only:
        if not args.xlsx:
            p.error("--spec-only requires --xlsx")
        _print(sync_spec_only(token, args.spec_only, args.xlsx, rates), ["spec_verification"])
        return
    if not args.json:
        p.error("--json is required (unless one of the *-only / --init-registry modes is used)")

    with open(args.json, "r", encoding="utf-8") as f:
        flat_data = json.load(f)
    result = upsert(token, flat_data, spec_xlsx_path=args.xlsx, fallback_id=args.fallback_id, rates=rates)
    if args.supplier_files_dir:
        paths = _docx_in(args.supplier_files_dir)
        if paths:
            result["supplier_files"] = attach_files(token, result["record_id"], paths)
    _print(result, ["verification", "spec_verification", "supplier_files"])


if __name__ == "__main__":
    try:
        main()
    except RuntimeError as e:  # API/config errors: one clean JSON line instead of a traceback
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)
