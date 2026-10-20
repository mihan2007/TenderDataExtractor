#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_supplier_files.py - one Word file per manufacturer from Items_and_prices.xlsx,
ready to send to suppliers as a request for quotation (step 7.1b in CLAUDE.md).

Each file is named after the brand (column "Manufacturer"; text in
parentheses is dropped: "Phoenix Contact (equivalent by specs)" ->
"Phoenix Contact.docx") and contains a short greeting and a table:
item no. in the tender | item name (with brand/model) | quantity (with unit).
The buyer's name and the tender's price ceilings are NOT included - the
files go to third parties. Rows without a manufacturer are listed under
"skipped" in the JSON output.

    python3 make_supplier_files.py --xlsx ".../AIExtract/Items_and_prices.xlsx" \
        --out-dir ".../AIExtract/Supplier requests" [--ref "T-12"] [--greeting "..."]

Output: {"status": "ok", "files": [{"brand", "path", "positions"}], "skipped": [...]}
Requires python-docx and openpyxl.
"""
import argparse
import json
import os
import re
import sys

try:
    import openpyxl
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as e:  # pragma: no cover
    print(json.dumps({"status": "error", "error": f"missing package: {e}. "
                      f"Install: {sys.executable} -m pip install python-docx openpyxl"}))
    sys.exit(1)

BAD_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*]+')
GREETING = ("Hello. We are preparing a purchase and would appreciate your quotation "
            "(price and lead time) for the items below.")
COL_WIDTHS = (Cm(2.6), Cm(11.4), Cm(3.0))  # 17 cm = A4 text width with 2 cm margins


def brand_of(manufacturer):
    brand = re.sub(r"\s*\([^)]*\)", "", str(manufacturer or "")).strip()
    return brand or str(manufacturer or "").strip()


def safe_filename(name):
    return BAD_FILENAME_CHARS.sub("_", name).strip(" .") or "Unnamed"


def fmt_qty(qty, unit):
    if qty is None or str(qty).strip() == "":
        return ""
    if isinstance(qty, float) and qty.is_integer():
        qty = int(qty)
    return f"{qty} {unit}".strip() if unit else str(qty)


def item_name(row):
    name = str(row.get("Item name") or "").strip()
    model = str(row.get("Brand/model") or "").strip()
    model = re.sub(r"^equivalent by specs:\s*", "", model, flags=re.IGNORECASE)
    if not model or model.startswith("—") or model.lower() in name.lower():
        return name
    return f"{name} — {model}"


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def build_doc(brand, rows, ref, path, greeting=GREETING):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(11)

    title = doc.add_paragraph().add_run(brand)
    title.bold = True
    title.font.size = Pt(16)
    sub = doc.add_paragraph().add_run("Request for price and lead time" + (f" (ref. {ref})" if ref else ""))
    sub.font.size = Pt(11)
    sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if greeting:
        g = doc.add_paragraph()
        g.paragraph_format.space_before = Pt(6)
        g.paragraph_format.space_after = Pt(10)
        g.add_run(greeting)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate(COL_WIDTHS):
        table.columns[i].width = w
    for i, h in enumerate(("Item no.", "Item", "Qty")):
        cell = table.rows[0].cells[i]
        cell.width = COL_WIDTHS[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
        shade(cell, "D9E2F3")
    for row in rows:
        cells = table.add_row().cells
        values = (str(row.get("No.") or ""), item_name(row), fmt_qty(row.get("Qty"), row.get("Unit")))
        for i, v in enumerate(values):
            cells[i].width = COL_WIDTHS[i]
            cells[i].text = v
    hdr = OxmlElement("w:tblHeader")  # repeat the header row on every page
    hdr.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(hdr)
    doc.save(path)


def main():
    ap = argparse.ArgumentParser(description="Per-manufacturer RFQ files from Items_and_prices.xlsx")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--greeting", default=GREETING, help="Text printed above the table")
    ap.add_argument("--ref", default="", help="Reference shown in the subtitle, e.g. the internal tender number")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    all_rows = list(wb.worksheets[0].iter_rows(values_only=True))
    headers = all_rows[0]
    if "Manufacturer" not in headers:
        print(json.dumps({"status": "error", "error": "the xlsx has no 'Manufacturer' column (step 5.1)"}))
        sys.exit(1)

    groups, order, skipped = {}, [], []
    for values in all_rows[1:]:
        row = dict(zip(headers, values))
        if not str(row.get("Item name") or "").strip():
            continue
        brand = brand_of(row.get("Manufacturer"))
        if not brand:
            skipped.append({"No.": row.get("No."), "item": row.get("Item name"), "reason": "no manufacturer"})
            continue
        if brand not in groups:
            groups[brand] = []
            order.append(brand)
        groups[brand].append(row)

    os.makedirs(args.out_dir, exist_ok=True)
    files = []
    for brand in order:
        path = os.path.join(args.out_dir, safe_filename(brand) + ".docx")
        build_doc(brand, groups[brand], args.ref, path, args.greeting)
        files.append({"brand": brand, "path": path, "positions": [r.get("No.") for r in groups[brand]]})
    print(json.dumps({"status": "ok", "files": files, "skipped": skipped}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
