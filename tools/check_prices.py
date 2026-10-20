#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
check_prices.py - checks that the market-price search in Items_and_prices.xlsx
is complete (steps 5.1g / 5.1i in CLAUDE.md). Read-only.

Run it BEFORE syncing the specification to NocoDB (step 7.1) and again after
the re-check pass (step 5.1i).

For every item with a brand/model (or an equivalent picked by specs) and for
each market (Domestic, International):
  * prices are numbers > 0 in ascending order, each with an https:// link,
    links are distinct;
  * if fewer than three prices were found, Notes must contain a marker like
        Domestic: found 1 of 3 [queries: "A" | "B" | "C" | "D" | "E"]
        International: not found [queries: "A" | "B" | "C" | "D" | "E" | "F"]
    listing at least --min-queries distinct queries (default 5);
  * generalizations instead of a search ("price on request from dealers",
    "no public prices for this brand") are flagged for incomplete markets.

Exit code: 0 - OK; 1 - problems found (go back to step 5.1i); 2 - read error.
--json <file> writes the list of items to re-check.

    python3 tools/check_prices.py --xlsx "<...>/AIExtract/Items_and_prices.xlsx" --json recheck.json
"""
import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from save_tender import parse_number_token  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("The openpyxl package is required: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

SHEET = "Items & prices"
MARKETS = ["Domestic", "International"]
NO_BRAND_VALUES = {"", "could not be determined", "not specified in the documents", "—", "-"}
SKIP_MARKERS = ("skipped (preliminary analysis)",)
NOTE_RE = re.compile(
    r"(?P<market>Domestic|International)\s*:\s*(?P<status>found\s*(?P<n>\d)\s*of\s*3|not\s+found)"
    r"\s*\[\s*queries\s*:\s*(?P<queries>[^\]]*)\]",
    re.IGNORECASE,
)
GENERALIZATION_RE = re.compile(
    r"price on request from (dealers|sellers|distributors)|no public prices|prices not publicly available|"
    r"brand .{0,30}(does not publish|has no) (public )?prices",
    re.IGNORECASE,
)


CURRENCY_RE = re.compile(r"[€$£¥₹₺]|\b[A-Z]{3}\b")


def to_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    # A foreign price may be written with its currency ("$45.90", "120 EUR") -
    # nocodb_client.py converts it later; here only the number matters.
    s = CURRENCY_RE.sub("", str(v)).strip()
    if not re.fullmatch(r"[\d\s.,'\u00a0]+", s):
        return "bad"
    value = parse_number_token(s)
    return "bad" if value is None else value


def parse_queries(text):
    qs = re.findall(r'["“]([^"”]+)["”]', text)
    return {q.strip().lower() for q in qs if q.strip()}


def check(path, min_queries):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    need = ["No.", "Brand/model", "Notes"]
    for m in MARKETS:
        for k in (1, 2, 3):
            need += [f"{m} offer {k}", f"{m} link {k}"]
    missing = [h for h in need if h not in idx]
    if missing:
        print("ERROR: missing columns: " + "; ".join(missing))
        return None, None

    problems, recheck = [], []
    total = checked = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        num = row[idx["No."]]
        if num is None or str(num).strip() == "":
            continue
        total += 1
        brand = str(row[idx["Brand/model"]] or "").strip()
        note = str(row[idx["Notes"]] or "")
        if brand.lower() in NO_BRAND_VALUES or any(s in note.lower() for s in SKIP_MARKERS):
            continue
        checked += 1
        notes = {mt.group("market").capitalize(): mt for mt in NOTE_RE.finditer(note)}
        row_probs, weak_markets = [], []
        for m in MARKETS:
            prices, links = [], []
            for k in (1, 2, 3):
                p = to_number(row[idx[f"{m} offer {k}"]])
                link = str(row[idx[f"{m} link {k}"]] or "").strip()
                if p == "bad":
                    row_probs.append(f"{m}: 'offer {k}' is not a number")
                    continue
                if p is None and not link:
                    continue
                if p is None or p <= 0:
                    row_probs.append(f"{m}: link {k} without a price")
                    continue
                if not link.startswith("https://"):
                    row_probs.append(f"{m}: price {k} has no https:// link")
                prices.append(p)
                links.append(link)
            if prices != sorted(prices):
                row_probs.append(f"{m}: prices are not in ascending order")
            if len(set(links)) != len(links):
                row_probs.append(f"{m}: duplicate links")
            if len(prices) < 3:
                weak_markets.append(m)
                mt = notes.get(m)
                if mt is None:
                    row_probs.append(f"{m}: found {len(prices)} of 3, but Notes has no "
                                     f"'{m}: ... [queries: \"...\" | ...]' marker")
                else:
                    nq = len(parse_queries(mt.group("queries")))
                    if nq < min_queries:
                        row_probs.append(f"{m}: found {len(prices)} of 3, but only {nq} queries (need >= {min_queries})")
        comments = " ".join(str(row[i] or "") for i, h in enumerate(header) if h.startswith("Agent comment"))
        if (GENERALIZATION_RE.search(note) or GENERALIZATION_RE.search(comments)) and \
                any(p.startswith(tuple(weak_markets)) for p in row_probs):
            row_probs.append("generalization instead of a search ('price on request from dealers' etc.) - "
                             "run the full query set")
        if row_probs:
            problems.append((num, brand, row_probs))
        if weak_markets:
            recheck.append({"No.": str(num), "Brand/model": brand, "markets": weak_markets})
    return (total, checked, problems), recheck


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, help="path to Items_and_prices.xlsx")
    ap.add_argument("--min-queries", type=int, default=5, help="minimum distinct queries per incomplete market")
    ap.add_argument("--json", help="write the items with < 3 prices here (input for step 5.1i)")
    a = ap.parse_args()
    try:
        res, recheck = check(a.xlsx, a.min_queries)
    except Exception as e:  # noqa: BLE001
        print(f"ERROR reading {a.xlsx}: {e}")
        sys.exit(2)
    if res is None:
        sys.exit(2)
    total, checked, problems = res
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(recheck, f, ensure_ascii=False, indent=1)
    print(f"Items: {total}, checked (with brand/equivalent): {checked}; markets: {', '.join(MARKETS)}")
    print(f"Items with < 3 prices in some market: {len(recheck)}")
    if not problems:
        print("OK: no problems")
        sys.exit(0)
    print(f"PROBLEMS - {len(problems)} item(s) (go back to step 5.1i and complete the search):")
    for num, brand, probs in problems:
        print(f"  No. {num} {brand[:60]}")
        for p in probs:
            print(f"     - {p}")
    sys.exit(1)


if __name__ == "__main__":
    main()
