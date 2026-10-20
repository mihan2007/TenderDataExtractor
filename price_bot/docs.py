"""Reading files sent to the bot: invoices, quotations, price lists.

extract(filename, data) -> {"text": str, "images": [bytes...], "kind": str, "note": str}
- text PDF, Word (.docx), Excel (.xlsx), CSV/TXT -> text;
- scanned PDF (no text layer) -> the first pages as images;
- photo / picture (jpg, png, webp, ...) -> image.
Images are downscaled to MAX_SIDE on the longer side (less traffic and tokens).
"""
import io
import os
import re

MAX_TEXT = 15000
MAX_PAGES_AS_IMAGES = 3
MAX_SIDE = 1800
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".heic"}


class DocError(Exception):
    pass


def _shrink(img) -> bytes:
    from PIL import Image
    img = img.convert("RGB")
    w, h = img.size
    k = MAX_SIDE / max(w, h)
    if k < 1:
        img = img.resize((int(w * k), int(h * k)), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf.getvalue()


def image_bytes(data: bytes) -> bytes:
    from PIL import Image
    try:
        return _shrink(Image.open(io.BytesIO(data)))
    except Exception as e:
        raise DocError(f"could not open the image: {e}")


def _page_text(page) -> str:
    """Page text that keeps table rows readable: layout mode, column gaps
    (3+ spaces) -> " | ". Falls back to plain extraction."""
    try:
        raw = page.extract_text(extraction_mode="layout") or ""
        lines = [re.sub(r" {3,}", " | ", ln.strip()) for ln in raw.splitlines()]
        return "\n".join(ln for ln in lines if ln)
    except Exception:
        return page.extract_text() or ""


def _pdf(data: bytes):
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [_page_text(p) for p in reader.pages[:20]]
    except Exception as e:
        raise DocError(f"could not read the PDF: {e}")
    text = "\n".join(pages).strip()
    if len(text) >= 40 * max(1, min(len(pages), 2)):
        return {"text": text, "images": [], "kind": "PDF", "note": ""}
    # almost no text - a scan: send the pages as images
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    images = [_shrink(pdf[i].render(scale=2).to_pil()) for i in range(min(len(pdf), MAX_PAGES_AS_IMAGES))]
    note = f"scan, pages sent: {len(images)}" + (f" of {len(pdf)}" if len(pdf) > len(images) else "")
    return {"text": text, "images": images, "kind": "scanned PDF", "note": note}


def _docx(data: bytes):
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise DocError(f"could not read the Word file: {e}")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                v = " ".join(c.text.split())
                if not cells or cells[-1] != v:  # merged cells repeat their text
                    cells.append(v)
            parts.append(" | ".join(cells))
    return {"text": "\n".join(parts), "images": [], "kind": "Word", "note": ""}


def _xlsx(data: bytes):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise DocError(f"could not read the Excel file: {e}")
    parts = []
    for ws in wb.worksheets[:5]:
        parts.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True, max_row=500):
            vals = [str(v).strip() for v in row if v not in (None, "")]
            if vals:
                parts.append(" | ".join(vals))
    return {"text": "\n".join(parts), "images": [], "kind": "Excel", "note": ""}


def _plain(data: bytes):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return {"text": data.decode(enc), "images": [], "kind": "text", "note": ""}
        except UnicodeDecodeError:
            continue
    raise DocError("could not detect the file encoding")


def extract(filename: str, data: bytes) -> dict:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".pdf":
        res = _pdf(data)
    elif ext == ".docx":
        res = _docx(data)
    elif ext in (".xlsx", ".xlsm"):
        res = _xlsx(data)
    elif ext in (".csv", ".txt"):
        res = _plain(data)
    elif ext in IMAGE_EXT:
        res = {"text": "", "images": [image_bytes(data)], "kind": "photo", "note": ""}
    elif ext in (".doc", ".xls"):
        raise DocError("legacy .doc/.xls is not supported - save as .docx/.xlsx or PDF")
    else:
        raise DocError(f"format {ext or '(no extension)'} is not supported (PDF, Word, Excel, CSV, images)")
    if len(res["text"]) > MAX_TEXT:
        res["text"] = res["text"][:MAX_TEXT]
        res["note"] = (res["note"] + "; " if res["note"] else "") + f"text truncated to {MAX_TEXT} characters"
    if not res["text"].strip() and not res["images"]:
        raise DocError("no text found in the file")
    return res
